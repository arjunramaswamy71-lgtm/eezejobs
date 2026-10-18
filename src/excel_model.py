"""Builds the downloadable eezejobs Excel model (docs/data/eezejobs_model.xlsx).

Pick a job from the dropdown and enter your pay; formulas look up the job's AI exposure,
its tasks, safer nearby jobs, hiring in the field, and what $X in its AI companies did
compared with the S&P 500. Every number is a formula or a labeled input, so it all recalculates.
"""

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).parent.parent
SITE = ROOT / "docs" / "data"
RAW = ROOT / "data" / "raw"
START = "2022-11-30"

F = "Arial"
BLUE, GREEN, BLACK, GREY = "0000FF", "008000", "000000", "666666"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F3A68")
SOFT = PatternFill("solid", fgColor="EEF2FA")
THIN = Border(bottom=Side(style="thin", color="CCCCCC"))


def font(**kw):
    return Font(name=F, **kw)


def header(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build():
    jobs = json.loads((SITE / "jobs.json").read_text())
    comps = json.loads((SITE / "companies.json").read_text())
    names = {c["ticker"]: c["name"] for c in comps if c.get("ticker")}
    prices = pd.read_csv(RAW / "prices.csv", index_col=0, parse_dates=True)
    hiring = pd.read_csv(RAW / "indeed_sectors_us.csv", parse_dates=["date"])
    hiring = hiring[hiring["variable"] == "total postings"].pivot_table(
        index="date", columns="display_name", values="indeed_job_postings_index")

    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    jw = wb.create_sheet("Jobs")
    gw = wb.create_sheet("Stock growth")
    hw = wb.create_sheet("Hiring")
    aw = wb.create_sheet("About")

    # ---------- Jobs data ----------
    cols = ["SOC code", "Job", "AI exposure score (0-100)", "US workers", "Median pay ($)",
            "Tasks AI can do", "Tasks AI can help", "Tasks that need a person", "Hiring sector",
            "Safer job 1", "Safer job 2", "Safer job 3"] + [f"AI stock {i}" for i in range(1, 9)]
    header(jw, 1, cols, [10, 46, 12, 12, 12, 10, 10, 12, 24, 30, 30, 30] + [9] * 8)
    title = {j["soc"]: j["title"] for j in jobs}
    for r, j in enumerate(sorted(jobs, key=lambda j: j["title"]), 2):
        safer = [title.get(s[0], "") for s in j.get("safer", [])][:3]
        safer += [""] * (3 - len(safer))
        basket = (j.get("basket") or [])[:8]
        basket += [""] * (8 - len(basket))
        row = [j["soc"], j["title"], j["score"], j["workers"], j["median_pay"],
               j["tasks"]["can"], j["tasks"]["helps"], j["tasks"]["you"], j.get("sector") or ""] + safer + basket
        for ci, v in enumerate(row, 1):
            c = jw.cell(row=r, column=ci, value=v)
            c.font = font(color=BLUE)
        jw.cell(row=r, column=4).number_format = "#,##0"
        jw.cell(row=r, column=5).number_format = "$#,##0"
    n_jobs = len(jobs) + 1
    jw.freeze_panes = "C2"
    jw.cell(row=n_jobs + 2, column=1, value="Source: eezejobs build of O*NET 31.0, BLS OEWS May 2025, Anthropic Economic Index, "
            "OpenAI GPTs are GPTs, Microsoft Working with AI, Felten-Raj-Seamans AIOE, and our 5-AI panel.").font = font(italic=True, color=GREY)

    # ---------- Stock growth ----------
    header(gw, 1, ["Ticker", "Company", f"Price on {START} ($)", "Latest price ($)", "Growth multiple", "Change"],
           [9, 26, 18, 16, 15, 11])
    p0 = prices[prices.index >= START].iloc[0]
    p1 = prices.ffill().iloc[-1]
    tick = sorted([t for t in prices.columns if t != "SPY"]) + ["SPY"]
    for r, t in enumerate(tick, 2):
        gw.cell(row=r, column=1, value=t).font = font(color=BLUE)
        gw.cell(row=r, column=2, value="S&P 500 index fund" if t == "SPY" else names.get(t, t)).font = font(color=BLUE)
        for ci, v in [(3, round(float(p0[t]), 2)), (4, round(float(p1[t]), 2))]:
            c = gw.cell(row=r, column=ci, value=v)
            c.font = font(color=BLUE)
            c.number_format = "$#,##0.00"
        c = gw.cell(row=r, column=5, value=f"=IF(C{r}>0,D{r}/C{r},\"\")")
        c.font = font()
        c.number_format = "0.00x"
        c = gw.cell(row=r, column=6, value=f"=IF(C{r}>0,D{r}/C{r}-1,\"\")")
        c.font = font()
        c.number_format = "0.0%;(0.0%);-"
    n_g = len(tick) + 1
    gw.cell(row=n_g + 2, column=1, value=f"Source: Yahoo Finance adjusted closing prices. Latest price as of {prices.index[-1]:%B %d, %Y}.").font = font(italic=True, color=GREY)

    # ---------- Hiring ----------
    header(hw, 1, ["Hiring sector (Indeed)", f"Job postings index, {START[:7]}", "Latest index", "Change"], [30, 22, 14, 11])
    h0 = hiring[hiring.index >= START].iloc[0]
    h1 = hiring.ffill().iloc[-1]
    for r, sec in enumerate(sorted(hiring.columns), 2):
        hw.cell(row=r, column=1, value=sec).font = font(color=BLUE)
        for ci, v in [(2, round(float(h0[sec]), 1)), (3, round(float(h1[sec]), 1))]:
            hw.cell(row=r, column=ci, value=v).font = font(color=BLUE)
        c = hw.cell(row=r, column=4, value=f"=IF(B{r}>0,C{r}/B{r}-1,\"\")")
        c.font = font()
        c.number_format = "0.0%;(0.0%);-"
    n_h = len(hiring.columns) + 1
    hw.cell(row=n_h + 2, column=1, value=f"Source: Indeed Hiring Lab job postings index (Feb 1, 2020 = 100), latest {hiring.index[-1]:%B %d, %Y}. CC-BY 4.0.").font = font(italic=True, color=GREY)

    # ---------- Start here ----------
    ws = start
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 54
    ws["A1"] = "eezejobs: your job and AI"
    ws["A1"].font = font(bold=True, size=18)
    ws["A2"] = "Change the yellow cells. Everything else updates by itself. For learning only, not financial or career advice."
    ws["A2"].font = font(italic=True, color=GREY)

    def label(r, text, bold=False):
        c = ws.cell(row=r, column=1, value=text)
        c.font = font(bold=bold)
        return c

    def note(r, text):
        ws.cell(row=r, column=3, value=text).font = font(color=GREY, size=9)

    def out(r, formula, fmt=None, green=False):
        c = ws.cell(row=r, column=2, value=formula)
        c.font = font(color=GREEN if green else BLACK, bold=False)
        if fmt:
            c.number_format = fmt
        c.border = THIN
        return c

    label(4, "Your inputs", True)
    jobs_range = f"Jobs!$B$2:$B${n_jobs}"
    label(5, "Your job")
    ws["B5"] = "Accountants and Auditors"
    label(6, "Your yearly pay ($)")
    ws["B6"] = 65000
    label(7, "Example amount to compare ($)")
    ws["B7"] = 1000
    for cell, fmt in [("B5", None), ("B6", "$#,##0"), ("B7", "$#,##0")]:
        ws[cell].font = font(color=BLUE)
        ws[cell].fill = YELLOW
        if fmt:
            ws[cell].number_format = fmt
    note(5, "Pick from the dropdown (example shown).")
    note(6, "Example value. Type your own.")
    note(7, "Example value. Used only for the comparison below.")
    dv = DataValidation(type="list", formula1=f"={jobs_range}", allow_blank=False)
    dv.error, dv.errorTitle = "Pick a job from the list.", "Job not found"
    ws.add_data_validation(dv)
    dv.add("B5")

    # Hidden helper: the job's row in the Jobs sheet
    ws["E5"] = f'=MATCH($B$5,{jobs_range},0)'
    ws["E5"].font = font(color=GREY, size=8)
    ws["D5"] = "row"
    ws["D5"].font = font(color=GREY, size=8)
    ws.column_dimensions["D"].hidden = True
    ws.column_dimensions["E"].hidden = True

    def J(col):
        return f"INDEX(Jobs!${col}$2:${col}${n_jobs},$E$5)"

    label(9, "Your job and AI", True)
    label(10, "AI exposure score (0 to 100)")
    out(10, f"=IFERROR({J('C')},\"\")", "0", True)
    note(10, "Higher than this share of all US jobs. Average of 6 sources.")
    label(11, "Exposure level")
    out(11, '=IF(B10="","",IF(B10>=75,"High",IF(B10>=45,"Medium","Low")))')
    label(12, "Tasks AI can do")
    out(12, f"=IFERROR({J('F')},\"\")", "0", True)
    label(13, "Tasks AI can help with")
    out(13, f"=IFERROR({J('G')},\"\")", "0", True)
    label(14, "Tasks that still need a person")
    out(14, f"=IFERROR({J('H')},\"\")", "0", True)
    label(15, "Share of tasks AI can do or help with")
    out(15, "=IFERROR((B12+B13)/(B12+B13+B14),\"\")", "0%")
    label(16, "US workers in this job")
    out(16, f"=IFERROR({J('D')},\"\")", "#,##0", True)
    label(17, "Typical (median) pay ($)")
    out(17, f"=IFERROR({J('E')},\"\")", "$#,##0", True)
    label(18, "Your pay vs typical")
    out(18, '=IFERROR(B6/B17-1,"")', "0.0%;(0.0%);-")
    label(19, "Pay linked to AI-exposed tasks ($, rough)")
    out(19, '=IFERROR(B6*B15,"")', "$#,##0")
    note(19, "Your pay times the share of tasks AI can do or help with. A rough way to size it, not a forecast of lost pay.")

    label(21, "Hiring in your field", True)
    label(22, "Indeed hiring sector")
    out(22, f'=IFERROR(IF({J("I")}="","Not matched",{J("I")}),"")', None, True)
    label(23, "Change in job postings since Nov 2022")
    out(23, f'=IFERROR(INDEX(Hiring!$D$2:$D${n_h},MATCH(B22,Hiring!$A$2:$A${n_h},0)),"n/a")', "0.0%;(0.0%);-", True)

    label(25, "Safer jobs near yours", True)
    note(25, "Similar skills, same or less training, similar pay, less AI exposure.")
    for k, col in enumerate(["J", "K", "L"]):
        r = 26 + k
        label(r, f"Option {k + 1}")
        out(r, f'=IFERROR(IF({J(col)}="","None found",{J(col)}),"")', None, True)
        c = ws.cell(row=r, column=3, value=f'=IFERROR(IF(B{r}="None found","","Exposure "&TEXT(INDEX(Jobs!$C$2:$C${n_jobs},MATCH(B{r},{jobs_range},0)),"0")&", median pay "&TEXT(INDEX(Jobs!$E$2:$E${n_jobs},MATCH(B{r},{jobs_range},0)),"$#,##0")),"")')
        c.font = font(color=GREEN, size=9)

    label(30, "Would owning the AI companies have helped?", True)
    note(30, "Equal amounts in each company from Nov 30, 2022 to the latest price.")
    for k, col in enumerate(["M", "N", "O", "P", "Q", "R", "S", "T"]):
        r = 31 + k
        ws.cell(row=r, column=4, value=f'=IFERROR(IF({J(col)}="","",{J(col)}),"")').font = font(color=GREY, size=8)
        ws.cell(row=r, column=5, value=f'=IF(D{r}="","",IFERROR(INDEX(\'Stock growth\'!$E$2:$E${n_g},MATCH(D{r},\'Stock growth\'!$A$2:$A${n_g},0)),""))').font = font(color=GREY, size=8)
    label(31, "AI companies for this job")
    out(31, '=IF(D31="","None matched",D31' + "".join(f'&IF(D{r}="","",", "&D{r})' for r in range(32, 39)) + ')')
    label(32, "Your example amount in those companies now ($)")
    out(32, '=IF(COUNT(E31:E38)=0,"",B7*AVERAGE(E31:E38))', "$#,##0")
    label(33, "Same amount in the S&P 500 now ($)")
    out(33, f"=B7*INDEX('Stock growth'!$E$2:$E${n_g},MATCH(\"SPY\",'Stock growth'!$A$2:$A${n_g},0))", "$#,##0", True)
    label(34, "Difference: AI companies minus S&P 500 ($)")
    out(34, '=IF(B32="","",B32-B33)', "$#,##0;($#,##0);-")
    label(35, "Verdict")
    out(35, '=IF(B32="","No AI companies matched to this job.",IF(B32>B33,"The AI companies did better than the S&P 500.","A plain S&P 500 fund did better than the AI companies."))')
    for r in range(4, 36):
        ws.row_dimensions[r].height = 18
    ws["A37"] = "Past results don't predict the future. Nothing here is a recommendation to buy or sell anything."
    ws["A37"].font = font(italic=True, color=GREY)
    ws.freeze_panes = "A4"

    # ---------- About ----------
    aw.column_dimensions["A"].width = 120
    lines = [
        ("eezejobs model: how it works", True),
        ("Blue text = inputs or source data. Black = formulas. Green = formulas pulling from another sheet. Yellow = cells you can change.", False),
        ("", False),
        ("AI exposure score: average percentile rank across 6 sources: real Claude usage (Anthropic Economic Index), real Copilot usage", False),
        ("(Microsoft Working with AI), GPT-4 and human expert task ratings (OpenAI GPTs are GPTs), AI occupational exposure (Felten, Raj, Seamans),", False),
        ("and our own panel of 5 AI models (GPT, Gemini, Claude, Llama, DeepSeek) rating all 18,428 O*NET tasks with the same rubric.", False),
        ("Tasks: each O*NET task is labeled AI can do it (direct exposure), AI can help (needs tools or partial usage), or needs a person.", False),
        ("Safer jobs: O*NET 31.0 related occupations and knowledge profiles, same or lower job zone, 85%+ of the pay, 10+ points less exposure, 10,000+ workers.", False),
        ("Stock comparison: equal amounts in each matched AI company, held from Nov 30, 2022 to the latest price, vs the SPY S&P 500 fund. Adjusted prices include dividends.", False),
        ("Pay linked to AI-exposed tasks is a rough sizing (pay x share of exposed tasks). Exposure means AI can change the work, not that the job disappears.", False),
        ("", False),
        ("Built by Arjun Ramaswamy, Virginia Tech. Free educational research project. Not financial, investment or career advice.", False),
    ]
    for r, (t, b) in enumerate(lines, 1):
        aw.cell(row=r, column=1, value=t).font = font(bold=b, size=14 if b else 10)

    for sheet in [jw, gw, hw]:
        sheet.sheet_view.showGridLines = True
    ws.sheet_view.showGridLines = False
    out_path = SITE / "eezejobs_model.xlsx"
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    print("Saved", build())
